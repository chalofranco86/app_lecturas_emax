from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

COLUMNAS = [
    ("ID", "id"),
    ("Código de Inmueble", "codigo_tarjeta"),
    ("Ruta", "ruta"),
    ("Correlativo", "correlativo"),
    ("Nombre del Inmueble", "nombre_inmueble"),
    ("Dirección", "direccion"),
    ("Usuario", "nombre_usuario"),
    ("Lectura", "lectura"),
    ("Mes de Proceso", "mes_proceso"),
    ("Fecha de Lectura", "fecha_lectura"),
    ("Observación", "observacion"),
    ("Coordenada X", "coordenada_x"),
    ("Coordenada Y", "coordenada_y"),
]


def generar_excel_lecturas(lecturas, nombre_hoja="Lecturas"):
    """Genera un archivo Excel a partir de una lista de lecturas."""
    if not lecturas:
        return None

    wb = Workbook()
    ws = wb.active
    ws.title = nombre_hoja

    # Estilos
    header_fill = PatternFill(start_color="1E88E5", end_color="1E88E5", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Encabezados
    for col_idx, (titulo, _) in enumerate(COLUMNAS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=titulo)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # Filas de datos
    for row_idx, lectura in enumerate(lecturas, start=2):
        for col_idx, (_, campo) in enumerate(COLUMNAS, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=lectura.get(campo, ""))
            cell.border = thin_border

    # Autoajustar ancho de columnas
    for col_idx, (titulo, _) in enumerate(COLUMNAS, start=1):
        max_len = len(titulo)
        for row_idx in range(2, len(lecturas) + 2):
            valor = ws.cell(row=row_idx, column=col_idx).value
            if valor is not None:
                max_len = max(max_len, len(str(valor)))
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = max_len + 3

    # Congelar la fila de encabezados
    ws.freeze_panes = "A2"

    # Guardar en memoria y devolver los bytes
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output